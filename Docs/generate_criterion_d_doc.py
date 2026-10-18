"""Generate the Criterion D (Development) explanation as an editable Word document."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT_PATH = Path(__file__).parent / "Criterion_D_Development.docx"

INTRO = (
    "This part explains the main code of my Reconciliation App. The app compares a ledger with bank "
    "statements and shows which rows match. I used Python (FastAPI) for the backend and SQLite for "
    "the database."
)

SECTIONS = [
    (
        "Connecting the database",
        '''engine = create_engine(settings.database_url, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()''',
        "I used SQLite because it is only one file, so nothing extra has to be installed. The create_engine "
        "line connects the app to that file. SessionLocal makes a new session for every request, and get_db "
        "gives it to the function and closes it at the end. I also switched on foreign keys, so the database "
        "does not allow a row that points to something that does not exist.",
        "The testing strategy is very effective because if the connection was wrong, the register and login "
        "tests would fail straight away.",
    ),
    (
        "Creating the tables",
        '''class LedgerEntry(Base):
    __tablename__ = "ledger_entry"
    id: Mapped[int] = mapped_column(primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("user.id", ondelete="CASCADE"))
    txn_date: Mapped[str]
    description: Mapped[str]
    amount_paise: Mapped[int]''',
        "Each table is a Python class, so I do not have to write SQL by hand. The line "
        "Base.metadata.create_all() in main.py makes all the tables when the app starts. There are tables for "
        "users, sessions, transactions, match results and ledger entries. The user_id column links every ledger "
        "entry to its owner, so one user cannot mix data with another user.",
        "The testing strategy is very effective as it checks that a saved record can be read back again.",
    ),
    (
        "Money stored as paise",
        '''def parse_amount_to_paise(raw: str) -> int:
    cleaned = _THOUSANDS_RE.sub("", raw).strip()
    value = Decimal(cleaned)
    paise = (value * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return int(paise)''',
        "Computers cannot store numbers like 0.01 exactly when float is used, and this can make two amounts "
        "look different by mistake. So I store money as whole numbers in paise. For example, 17.70 rupees is "
        "saved as 1770. This method also removes commas, so 1,25,000.00 can be read.",
        "The testing strategy is very effective as the tolerance tests check that amounts with small "
        "differences are compared correctly.",
    ),
    (
        "Creating a ledger entry",
        '''@router.post("", response_model=LedgerEntryOut, status_code=201)
def create_entry(payload: LedgerEntryIn, current_user=Depends(get_current_user), db=Depends(get_db)):
    entry = LedgerEntry(user_id=current_user.id, txn_date=payload.txn_date,
        description=payload.description,
        amount_paise=parse_amount_to_paise(str(payload.amount)))
    db.add(entry)
    db.commit()
    db.refresh(entry)''',
        "This method adds a new row to the ledger. The data first goes through LedgerEntryIn, which checks that "
        "the description is not empty, the amount is more than zero and the date looks like YYYY-MM-DD. Then "
        "the amount is changed to paise, and db.add and db.commit save it. The current logged-in user is "
        "added automatically.",
        "The testing strategy is very effective as it checks that a new entry is saved and the same amount "
        "comes back when it is read.",
    ),
    (
        "Showing the ledger entries",
        '''entries = (db.query(LedgerEntry)
    .filter(LedgerEntry.user_id == current_user.id)
    .order_by(LedgerEntry.txn_date.desc(), LedgerEntry.id.desc())
    .all())''',
        "This query reads all the entries of the logged-in user only. The newest date is shown first. "
        "The Angular page calls this method and shows the result in a table, and the table is refreshed after "
        "every add, edit or delete.",
        "The testing strategy is very effective as it checks that the saved entry appears in the list.",
    ),
    (
        "Updating a ledger entry",
        '''entry = _get_owned_entry(entry_id, current_user, db)
entry.txn_date = payload.txn_date
entry.description = payload.description
entry.amount_paise = parse_amount_to_paise(str(payload.amount))
db.commit()''',
        "To change a row, I first find it using _get_owned_entry. This also checks that the row belongs to "
        "the user. Then each value is replaced with the new one and db.commit saves it. The same validation "
        "is used as in creating, so a bad value cannot be saved while editing.",
        "The testing strategy is very effective as it checks that the changed values appear after saving.",
    ),
    (
        "Deleting a ledger entry",
        '''def _get_owned_entry(entry_id, user, db):
    entry = db.get(LedgerEntry, entry_id)
    if entry is None or entry.user_id != user.id:
        raise HTTPException(status_code=404, detail="Ledger entry not found")
    return entry

db.delete(entry)
db.commit()''',
        "Before deleting, the code finds the entry and checks the owner. If the row is not found, or it is "
        "somebody else's row, the same 404 message is shown. This way a stranger cannot guess which ids "
        "exist. Then db.delete removes the row and db.commit makes it permanent.",
        "The testing strategy is very effective as it checks that a deleted entry is gone from the list.",
    ),
    (
        "Uploading files and creating a session",
        '''session = ReconciliationSession(user_id=current_user.id, ledger_filename=ledger_filename,
    bank_filename=bank_paths_and_labels[0][1],
    date_tolerance_days=date_tolerance_days, amount_tolerance=amount_tolerance,
    tolerance_type=tolerance_type, status="pending", progress_pct=0)
db.add(session)
db.commit()
start_reconciliation_job(session.id, ledger_path, bank_paths_and_labels)''',
        "The user can upload up to three bank statements. Every file is checked to be a real PDF before it "
        "is saved with a random name. The session row stores the settings, like date tolerance and amount "
        "tolerance. The ledger can come from a PDF or from the saved ledger entries. The matching work then "
        "starts in a separate thread, so the page does not freeze.",
        "The testing strategy is very effective as it checks that wrong files are refused.",
    ),
    (
        "Matching the records",
        '''for bank_row in unconsumed_bank:
    date_diff = _signed_days_between(ledger_row.txn_date, bank_row.txn_date)
    if abs(date_diff) > date_tolerance_days:
        continue
    amount_diff = bank_row.amount_paise - ledger_row.amount_paise
    if abs(amount_diff) > allowed:
        continue
    score = _similarity(ledger_row.description, bank_row.description)
    candidates.append((bank_row, date_diff, amount_diff, score))''',
        "This is the main part of the app. For each ledger row, I loop through the bank rows that are not "
        "used yet. A bank row is skipped if its date or amount is too far away. The rest are kept, and "
        "rapidfuzz gives each a score for how alike the descriptions are. The row with the best score is "
        "chosen and removed from the list, so one bank row is never used twice. If the score is 80 or more "
        "it is Matched, if it is lower it is Partial, and if nothing is left it is Missing.",
        "The testing strategy is very effective as I compared the result with a sample file I made by hand "
        "and all the counts were the same.",
    ),
    (
        "Finding duplicates",
        '''key = (row.txn_date, row.description.strip().lower(), row.amount_paise)
groups.setdefault(key, []).append(row)
for group in groups.values():
    if len(group) > 1:
        duplicates.update(r.row_number for r in group)''',
        "Rows with the same date, description and amount are put in one group. If a group has more than one "
        "row, all of them are marked as duplicates. I used a dictionary because it finds the groups in one "
        "pass. The description is changed to small letters so that capital letters do not matter.",
        "The testing strategy is very effective as it checks that two same rows are flagged.",
    ),
    (
        "Saving the results",
        '''db.add(MatchResult(session_id=session.id,
    ledger_txn_id=ledger_txn.id if ledger_txn else None,
    bank_txn_id=bank_txn.id if bank_txn else None,
    match_status=outcome.status, reason=outcome.reason,
    similarity_score=outcome.similarity_score))''',
        "After matching, every row is saved as a Txn and every result as a MatchResult. The ledger or bank id "
        "can be empty, because a Missing row only has one side. Progress goes up from 5 to 100 percent while "
        "this happens, and the Angular page asks for it again and again to show the progress bar. If "
        "something fails, the status is set to error with a message.",
        "The testing strategy is very effective as it checks that the saved totals match the expected ones.",
    ),
    (
        "Deleting a session",
        '''session = _get_owned_session(session_id, current_user, db)
db.delete(session)
db.commit()

transactions = relationship(back_populates="session", cascade="all, delete-orphan")''',
        "When a session is deleted from History, its transactions and results should go too. I used "
        "cascade=\"all, delete-orphan\", so SQLAlchemy removes the child rows by itself. This keeps the "
        "database clean and there are no leftover rows.",
        "The testing strategy does not specifically test this part, so I checked it by hand in the History page.",
    ),
    (
        "Updating an old database",
        '''_NEW_COLUMNS = {"txn": [("bank_label", "TEXT")]}

for column_name, column_type in columns:
    if existing and column_name not in existing:
        conn.exec_driver_sql(f"ALTER TABLE {table} ADD COLUMN {column_name} {column_type}")''',
        "When I added new fields, the old database file did not have them, and create_all does not change "
        "tables that already exist. So this small method checks the columns and adds the missing ones with "
        "ALTER TABLE. Because of this, old data is not lost.",
        "The testing strategy does not specifically test this part, so I checked it by hand with an older database file.",
    ),
]


def add_code(doc: Document, code: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Pt(12)
    for i, line in enumerate(code.split("\n")):
        run = para.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x26, 0x32, 0x38)
        if i < len(code.split("\n")) - 1:
            run.add_break()


def build() -> int:
    doc = Document()
    title = doc.add_heading("Criterion D - Development", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(INTRO)
    words = len(INTRO.split())

    for heading, code, explanation, test in SECTIONS:
        doc.add_heading(heading, level=2)
        add_code(doc, code)
        doc.add_paragraph(explanation)
        doc.add_paragraph(test)
        words += len((explanation + " " + test).split())

    doc.save(OUT_PATH)
    return words


if __name__ == "__main__":
    print(f"Saved {OUT_PATH} ({build()} words, code excluded)")
